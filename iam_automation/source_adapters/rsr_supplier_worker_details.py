"""
File: iam_automation/source_adapters/rsr_supplier_worker_details.py
Purpose: Provides the rsr supplier worker details component of IAM Automation.
Responsibilities: Exposes or coordinates: RsrIngestionResult, _utc_now, _header_key, _value, RsrSupplierWorkerDetailsAdapter.
Dependencies: openpyxl, sqlalchemy, infrastructure, application
Inputs/outputs: Inputs: application arguments, configured paths, and persisted runtime data. Outputs: domain results, state changes, or operational records.
Author / implementation: Miguel Fernandez.
Developed for: Kaiser Permanente NSES GWFM.
Project implementation and integration: Miguel Fernandez.
Third-party source, libraries, and dependencies retain their original licenses.
Estimated implementation effort: 1.25 h (Git-derived heuristic; see docs/file-effort-history.md).
"""

from __future__ import annotations

import hashlib
import json
import uuid
import warnings
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook
from sqlalchemy import func, select

from ..infrastructure.sqlite.orm import (
    ImportBatchModel,
    IngestionRecordModel,
    QuarantineRecordModel,
    WorkerDiscrepancyModel,
    WorkerEngagementModel,
    WorkerIdentifierModel,
    WorkerModel,
    WorkerSourceLineageModel,
    orm_session,
)
from ..application.stable_file_read import file_modified_at_utc, read_stable_bytes
from ..application.ingestion_dedup import latest_payload_digests, payload_digest, serialize_payload
from ..application.star_schema_writer import StarSchemaWriter
from ..application.worker_history_service import record_worker_version
from ..application.source_reconciliation import ensure_source_identifiers, find_worker_matches
from ..application.source_change_service import apply_engagement_changes, material_source_fields, record_changed_lineage


SOURCE_SYSTEM = "RSR_SUPPLIER_WORKER_DETAILS"
REQUIRED_HEADERS = {"supplier_reference", "worker_id", "worker"}

# Fieldglass exports use different labels for the same worker-level concepts
# depending on the report/filter set. Keep one internal vocabulary so both the
# classic RSR Supplier Worker Details report and the Headcount report follow
# the same approval and reconciliation path.
HEADER_ALIASES = {
    "contractor_id": "worker_id",
    "contractor": "worker",
    "contractor_status": "worker_status",
    "contractor_start_date": "worker_start_date",
    "contractor_end_date": "worker_end_date",
    "statement_of_work_id": "work_order_id",
    "external_reference_id": "statement_of_work",
    "contractor_owner": "worker_primary_contact",
    "supplier": "supplier_name",
}


@dataclass(frozen=True)
class RsrIngestionResult:
    batch_id: str
    filename: str
    total_rows: int
    staged_rows: int
    quarantined_rows: int
    already_processed: bool = False


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _header_key(value: object) -> str:
    return "_".join(str(value or "").strip().casefold().replace("/", " ").replace("-", " ").split())


def _value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value).strip()


class RsrSupplierWorkerDetailsAdapter:
    """Stage and reconcile RSR details without making the workbook canonical."""

    source_system = SOURCE_SYSTEM

    def __init__(self, database_path: Path):
        self.database_path = database_path

    def ingest_file(self, path: Path, actor: str = "workspace-watcher") -> RsrIngestionResult:
        path = path.expanduser().resolve()
        checksum = hashlib.sha256(read_stable_bytes(path)).hexdigest()
        with orm_session(self.database_path) as session:
            existing = session.scalar(select(IngestionRecordModel).where(IngestionRecordModel.source_file == str(path), IngestionRecordModel.source_checksum == checksum).limit(1))
            if existing:
                return RsrIngestionResult(existing.batch_id, path.name, 0, 0, 0, True)

        rows = self._read_rows(path)
        now = _utc_now()
        batch_id = str(uuid.uuid4())
        staged = quarantined = 0
        with orm_session(self.database_path) as session:
            batch = ImportBatchModel(batch_id=batch_id, filename=path.name, source_type=SOURCE_SYSTEM, status="PROCESSING", total_rows=len(rows), created_by=actor, created_at_utc=now, source_modified_at_utc=file_modified_at_utc(path))
            session.add(batch)
            existing_payloads = latest_payload_digests(session, SOURCE_SYSTEM, {self._source_key(payload, row_number) for row_number, payload in rows})
            for row_number, payload in rows:
                source_key = self._source_key(payload, row_number)
                normalized_json = serialize_payload(payload)
                record_id = str(uuid.uuid4())
                missing = sorted(REQUIRED_HEADERS - {key for key, value in payload.items() if value})
                if missing:
                    session.add(QuarantineRecordModel(quarantine_id=str(uuid.uuid4()), batch_id=batch_id, row_number=row_number, reason=f"Missing required fields: {', '.join(missing)}", payload_json=json.dumps(payload, sort_keys=True), status="OPEN", created_at_utc=now))
                    session.add(IngestionRecordModel(record_id=record_id, batch_id=batch_id, source_system=SOURCE_SYSTEM, source_file=str(path), source_checksum=checksum, row_number=row_number, source_record_key=source_key, normalized_json=json.dumps(payload, sort_keys=True), payload_json=json.dumps(payload, sort_keys=True), status="QUARANTINED", match_status="NOT_EVALUATED", match_reason="Missing required fields", created_at_utc=now))
                    quarantined += 1
                    continue
                if existing_payloads.get(source_key) == payload_digest(normalized_json):
                    continue
                session.add(IngestionRecordModel(record_id=record_id, batch_id=batch_id, source_system=SOURCE_SYSTEM, source_file=str(path), source_checksum=checksum, row_number=row_number, source_record_key=source_key, normalized_json=normalized_json, payload_json=json.dumps(payload, sort_keys=True), status="STAGED", match_status="NOT_EVALUATED", created_at_utc=now))
                staged += 1
            StarSchemaWriter.write_batch_in_session(session, batch_id, now)
            batch.valid_rows = staged
            batch.quarantine_rows = quarantined
            batch.status = "COMPLETED" if quarantined == 0 else "COMPLETED_WITH_WARNINGS"
            batch.completed_at_utc = _utc_now()
            session.commit()
        self.match_batch(batch_id)
        StarSchemaWriter.clear_legacy_payloads(self.database_path, batch_id)
        return RsrIngestionResult(batch_id, path.name, len(rows), staged, quarantined)

    def match_batch(self, batch_id: str) -> list[dict[str, str]]:
        with orm_session(self.database_path) as session:
            records = session.scalars(select(IngestionRecordModel).where(IngestionRecordModel.batch_id == batch_id, IngestionRecordModel.source_system == SOURCE_SYSTEM, IngestionRecordModel.status == "STAGED").order_by(IngestionRecordModel.row_number)).all()
            preview = []
            for record in records:
                payload = StarSchemaWriter.payload_for_record(session, record)
                matches, match_reason = self._find_match_detail(session, payload)
                if len(matches) == 1:
                    worker_id = matches[0]
                    conflicts = self._record_conflicts(session, worker_id, record, payload)
                    # A unique source identifier is authoritative for identity.
                    # Keep name differences as audit discrepancies, but do not
                    # block enrichment approval because of display-name drift.
                    status = "MATCHED"
                    reason = match_reason if not conflicts else f"{match_reason}; {len(conflicts)} field conflict(s) preserved for review"
                    record.matched_worker_id = worker_id
                elif len(matches) > 1:
                    status, reason = "AMBIGUOUS", "Identifiers resolve to multiple canonical workers"
                    record.matched_worker_id = None
                else:
                    status, reason = "UNMATCHED", "No exact canonical identifier match; eligible for approved new Employee creation"
                    record.matched_worker_id = None
                record.match_status = status
                record.match_reason = reason
                preview.append({"record_id": record.record_id, "source_record_key": record.source_record_key, "match_status": status, "matched_worker_id": record.matched_worker_id or "", "match_reason": reason})
            session.commit()
            return preview

    def refresh_matches(self) -> dict[str, int]:
        """Re-evaluate staged RSR rows after another source creates workers.

        RSR is an enrichment source, so a row can be unmatched when it first
        arrives and become eligible after an approved Roster Lock promotion.
        """
        matched = unmatched = ambiguous = conflicts = reviewed = vendor_repaired = 0
        with orm_session(self.database_path) as session:
            records = session.scalars(
                select(IngestionRecordModel)
                .where(
                    IngestionRecordModel.source_system == SOURCE_SYSTEM,
                    IngestionRecordModel.status == "STAGED",
                )
                .order_by(IngestionRecordModel.created_at_utc, IngestionRecordModel.row_number)
            ).all()
            for record in records:
                payload = StarSchemaWriter.payload_for_record(session, record)
                matches, match_reason = self._find_match_detail(session, payload)
                previous = record.match_status
                if len(matches) == 1:
                    worker_id = matches[0]
                    conflict_fields = self._record_conflicts(session, worker_id, record, payload)
                    record.matched_worker_id = worker_id
                    # The ID match is safe to approve; discrepancies remain
                    # visible in the audit/review views.
                    record.match_status = "MATCHED"
                    record.match_reason = match_reason if not conflict_fields else f"{match_reason}; {len(conflict_fields)} field conflict(s) preserved for review"
                elif len(matches) > 1:
                    record.matched_worker_id = None
                    record.match_status = "AMBIGUOUS"
                    record.match_reason = "Identifiers resolve to multiple canonical workers"
                else:
                    record.matched_worker_id = None
                    record.match_status = "UNMATCHED"
                    record.match_reason = "No exact canonical identifier match; eligible for approved new Employee creation"
                reviewed += int(previous != record.match_status)
                if record.match_status == "MATCHED":
                    matched += 1
                elif record.match_status == "AMBIGUOUS":
                    ambiguous += 1
                elif record.match_status == "CONFLICT":
                    conflicts += 1
                else:
                    unmatched += 1

            # Older applied rows were created before vendor derivation existed.
            # Repair only RSR engagements from their immutable source payload.
            applied_records = session.scalars(
                select(IngestionRecordModel).where(
                    IngestionRecordModel.source_system == SOURCE_SYSTEM,
                    IngestionRecordModel.status == "APPLIED",
                )
            ).all()
            source_by_worker_and_order: dict[tuple[str, str], dict[str, str]] = {}
            for record in applied_records:
                payload = StarSchemaWriter.payload_for_record(session, record)
                worker_id = record.matched_worker_id
                work_order_id = payload.get("work_order_id", "")
                if worker_id and work_order_id:
                    source_by_worker_and_order[(worker_id, work_order_id)] = payload
            engagements = session.scalars(
                select(WorkerEngagementModel).where(WorkerEngagementModel.source_system == SOURCE_SYSTEM)
            ).all()
            for engagement in engagements:
                payload = source_by_worker_and_order.get((engagement.worker_id, engagement.work_order_id or ""))
                if not payload:
                    continue
                vendor_name = self._derive_vendor_name(payload)
                if vendor_name and engagement.vendor_name != vendor_name:
                    engagement.vendor_name = vendor_name
                    vendor_repaired += 1
            session.commit()
        return {"reviewed": reviewed, "matched": matched, "unmatched": unmatched, "ambiguous": ambiguous, "conflicts": conflicts, "vendor_repaired": vendor_repaired}

    def apply_approved_records(
        self,
        record_ids: list[str],
        actor: str,
        *,
        progress_callback: Callable[[int, int, int, int, int], None] | None = None,
    ) -> dict[str, object]:
        selected = {str(value) for value in record_ids if str(value).strip()}
        if not selected:
            raise ValueError("The RSR import request contains no records.")
        now = _utc_now()
        applied = created = reused = changed = unchanged = failed = 0
        errors: list[dict[str, str]] = []
        with orm_session(self.database_path) as session:
            records = session.scalars(select(IngestionRecordModel).where(IngestionRecordModel.record_id.in_(selected), IngestionRecordModel.source_system == SOURCE_SYSTEM)).all()
            handled = 0
            for record in records:
                if record.status == "APPLIED":
                    reused += 1
                    handled += 1
                    if progress_callback:
                        progress_callback(handled, len(records), created, reused, failed)
                    continue
                if record.status != "STAGED" or record.match_status not in {"MATCHED", "UNMATCHED"}:
                    errors.append({"record_id": record.record_id, "error": f"Record is not eligible for approval: {record.status}/{record.match_status}."})
                    failed += 1
                    handled += 1
                    if progress_callback:
                        progress_callback(handled, len(records), created, reused, failed)
                    continue
                try:
                    payload = StarSchemaWriter.payload_for_record(session, record)
                    worker_id = record.matched_worker_id
                    if record.match_status == "UNMATCHED":
                        # Matching is calculated when the file is ingested, but
                        # another source may have created the canonical worker
                        # before approval. Revalidate inside this transaction
                        # before ever creating a new canonical employee.
                        matches, match_reason = find_worker_matches(session, SOURCE_SYSTEM, payload)
                        if len(matches) > 1:
                            raise ValueError("The RSR record matches multiple canonical employees and requires manual resolution.")
                        if len(matches) == 1 and match_reason != "Canonical name match":
                            worker_id = matches[0]
                            record.matched_worker_id = worker_id
                            record.match_status = "MATCHED"
                            record.match_reason = "Canonical match revalidated during approval"
                    was_created = worker_id is None
                    if record.match_status == "UNMATCHED":
                        worker_id = str(uuid.uuid4())
                        first_name = str(payload.get("first_name") or "").strip()
                        last_name = str(payload.get("last_name") or "").strip()
                        session.add(WorkerModel(
                            worker_id=worker_id,
                            first_name=first_name,
                            last_name=last_name,
                            display_name=" ".join(part for part in (first_name, last_name) if part),
                            canonical_email=str(payload.get("email") or "").strip().lower(),
                            # Imported workers enter the canonical directory as
                            # active; later status changes use the audited
                            # employee request workflow.
                            status="ACTIVE",
                            record_version=1,
                            created_at_utc=now,
                            updated_at_utc=now,
                        ))
                        created += 1
                    else:
                        reused += 1
                    source_initial = False
                    engagement = session.scalar(select(WorkerEngagementModel).where(WorkerEngagementModel.worker_id == worker_id, WorkerEngagementModel.source_system == SOURCE_SYSTEM, WorkerEngagementModel.work_order_id == payload.get("work_order_id", "")).order_by(WorkerEngagementModel.updated_at_utc.desc()).limit(1))
                    if was_created:
                        engagement = WorkerEngagementModel(engagement_id=str(uuid.uuid4()), worker_id=worker_id, source_system=SOURCE_SYSTEM, created_at_utc=now, updated_at_utc=now)
                        session.add(engagement)
                        self._apply_engagement(engagement, payload, now)
                        changed_fields = set(payload)
                    else:
                        changed_fields = material_source_fields(session, worker_id, SOURCE_SYSTEM, payload)
                        if changed_fields - {"identifiers"}:
                            if engagement is None:
                                source_initial = True
                                engagement = WorkerEngagementModel(engagement_id=str(uuid.uuid4()), worker_id=worker_id, source_system=SOURCE_SYSTEM, created_at_utc=now, updated_at_utc=now)
                                session.add(engagement)
                                self._apply_engagement(engagement, payload, now)
                            else:
                                apply_engagement_changes(session, engagement, payload, changed_fields, now)
                    self._add_identifiers(session, worker_id, payload, now)
                    if was_created or source_initial:
                        for field_name, value in payload.items():
                            session.add(WorkerSourceLineageModel(lineage_id=str(uuid.uuid4()), worker_id=worker_id, engagement_id=engagement.engagement_id, batch_id=record.batch_id, source_system=SOURCE_SYSTEM, source_file=record.source_file, source_checksum=record.source_checksum, row_number=record.row_number, field_name=field_name, raw_value=str(value or ""), normalized_value=str(value or ""), observed_at_utc=now))
                    elif changed_fields - {"identifiers"}:
                        record_changed_lineage(session, record, worker_id, now, changed_fields, engagement.engagement_id if engagement else None)
                    if changed_fields:
                        record_worker_version(session, worker_id, actor=actor, change_type="RSR_INITIAL_PROFILE" if was_created else "RSR_SOURCE_ENRICHMENT", source_system=record.source_system, source_file=record.source_file, source_checksum=record.source_checksum, source_record_key=record.source_record_key, changed_fields=sorted(changed_fields))
                    if not was_created:
                        changed += int(bool(changed_fields))
                        unchanged += int(not changed_fields)
                    record.matched_worker_id = worker_id
                    record.match_status = "MATCHED"
                    if was_created:
                        record.match_reason = "Canonical employee created during approval"
                    record.status = "APPLIED"
                    record.observation_state = "CURRENT"
                    record.linked_at_utc = record.linked_at_utc or now
                    record.last_seen_at_utc = now
                    applied += 1
                except (ValueError, json.JSONDecodeError) as exc:
                    errors.append({"record_id": record.record_id, "error": str(exc)})
                    failed += 1
                handled += 1
                if progress_callback:
                    progress_callback(handled, len(records), created, reused, failed)
            session.commit()
        return {"processed": applied, "applied": applied, "created": created, "reused": reused, "changed": changed, "unchanged": unchanged, "failed": failed, "errors": errors, "actor": actor}

    @staticmethod
    def _read_rows(path: Path) -> list[tuple[int, dict[str, str]]]:
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore",
                message="Data Validation extension is not supported and will be removed",
                category=UserWarning,
            )
            workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            header_row = None
            headers: list[str] = []
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                normalized = [HEADER_ALIASES.get(_header_key(value), _header_key(value)) for value in row]
                if REQUIRED_HEADERS.issubset(set(normalized)):
                    header_row, headers = row_number, normalized
                    break
            if header_row is None:
                raise ValueError("RSR workbook does not contain the expected Worker Details header row.")
            result = []
            for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
                if not any(value not in (None, "") for value in row):
                    continue
                payload = {headers[index]: _value(row[index] if index < len(row) else None) for index in range(len(headers)) if headers[index]}
                payload = RsrSupplierWorkerDetailsAdapter._normalize_payload(payload)
                if RsrSupplierWorkerDetailsAdapter._is_report_footer(payload):
                    continue
                result.append((row_number, payload))
            return result
        finally:
            workbook.close()

    @staticmethod
    def _normalize_payload(payload: dict[str, str]) -> dict[str, str]:
        worker = payload.get("worker", "")
        if "," in worker:
            last, first = (part.strip() for part in worker.split(",", 1))
            payload.setdefault("first_name", first)
            payload.setdefault("last_name", last)
        else:
            parts = worker.split()
            payload.setdefault("first_name", parts[0] if parts else "")
            payload.setdefault("last_name", " ".join(parts[1:]))
        payload["supplier_reference"] = payload.get("supplier_reference", "").strip()
        payload["worker_id"] = payload.get("worker_id", "").strip()
        return payload

    @staticmethod
    def _is_report_footer(payload: dict[str, str]) -> bool:
        """Ignore Fieldglass report metadata rows after the worker data."""
        if any(str(payload.get(field) or "").strip() for field in REQUIRED_HEADERS):
            return False
        report_text = " ".join(str(value or "").strip().casefold() for value in payload.values())
        return (
            "the data contained in this report is as of" in report_text
            or "date:" in report_text and "us/central" in report_text
        )

    @staticmethod
    def _derive_vendor_name(payload: dict[str, str]) -> str:
        """Return the business vendor, keeping the Fieldglass SOW intact."""
        for key in ("vendor_name", "vendor", "supplier_name", "supplier"):
            value = str(payload.get(key) or "").strip()
            if value:
                return RsrSupplierWorkerDetailsAdapter._canonical_vendor_name(value)
        sow = str(payload.get("statement_of_work") or "").strip()
        if sow:
            return RsrSupplierWorkerDetailsAdapter._canonical_vendor_name(sow.split(maxsplit=1)[0])
        return "Fieldglass"

    @staticmethod
    def _canonical_vendor_name(value: str) -> str:
        """Collapse known Fieldglass vendor aliases without changing source data."""
        normalized = value.strip()
        if normalized.casefold() in {"ttec_2020", "ttec-2020", "ttec 2020"}:
            return "TTEC"
        return normalized

    @staticmethod
    def _source_key(payload: dict[str, str], row_number: int) -> str:
        return f"{payload.get('supplier_reference') or payload.get('worker_id') or 'ROW'}:{row_number}"

    @staticmethod
    def _find_match_detail(session, payload: dict[str, str]) -> tuple[list[str], str]:
        return find_worker_matches(session, SOURCE_SYSTEM, payload)

    @staticmethod
    def _find_matches(session, payload: dict[str, str]) -> list[str]:
        return RsrSupplierWorkerDetailsAdapter._find_match_detail(session, payload)[0]

    @staticmethod
    def _record_conflicts(session, worker_id: str, record: IngestionRecordModel, payload: dict[str, str]) -> list[str]:
        worker = session.get(WorkerModel, worker_id)
        if worker is None:
            return ["worker"]
        current = {"first_name": worker.first_name, "last_name": worker.last_name, "email": worker.canonical_email}
        conflicts = []
        for field_name, incoming in current.items():
            value = str(payload.get(field_name, "")).strip()
            if value and incoming and value.casefold() != str(incoming).casefold():
                conflicts.append(field_name)
                existing = session.scalar(select(WorkerDiscrepancyModel).where(
                    WorkerDiscrepancyModel.worker_id == worker_id,
                    WorkerDiscrepancyModel.source_system == SOURCE_SYSTEM,
                    WorkerDiscrepancyModel.source_record_key == record.source_record_key,
                    WorkerDiscrepancyModel.field_name == field_name,
                    WorkerDiscrepancyModel.status == "OPEN",
                ).limit(1))
                if existing is None:
                    session.add(WorkerDiscrepancyModel(discrepancy_id=str(uuid.uuid4()), worker_id=worker_id, source_system=SOURCE_SYSTEM, source_record_key=record.source_record_key, field_name=field_name, current_value=str(incoming), incoming_value=value, status="OPEN", created_at_utc=_utc_now()))
        return conflicts

    @staticmethod
    def _apply_engagement(engagement: WorkerEngagementModel, payload: dict[str, str], now: str) -> None:
        mapping = {"statement_of_work": "statement_of_work", "work_order_id": "work_order_id", "worker_primary_contact": "primary_contact_name", "lob": "lob", "sow_worker_role": "job_title", "worker_start_date": "onboarding_date", "worker_end_date": "production_start_date", "resident_license_number": "license_number", "resident_license_state": "license_state", "npn": "npn", "taxonomy": "taxonomy_code"}
        for source, target in mapping.items():
            value = str(payload.get(source, "")).strip()
            if value:
                setattr(engagement, target, value)
        engagement.vendor_name = RsrSupplierWorkerDetailsAdapter._derive_vendor_name(payload)
        engagement.status = payload.get("worker_status") or engagement.status or "PENDING"
        engagement.updated_at_utc = now

    @staticmethod
    def _add_identifiers(session, worker_id: str, payload: dict[str, str], now: str) -> None:
        ensure_source_identifiers(session, worker_id, SOURCE_SYSTEM, payload, now)
#endfile
