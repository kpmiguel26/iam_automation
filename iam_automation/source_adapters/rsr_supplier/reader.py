"""Read Fieldglass RSR Supplier Worker Details workbooks."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .normalizer import HEADER_ALIASES, header_key, is_report_footer, normalize_payload, value


REQUIRED_HEADERS = {"supplier_reference", "worker_id", "worker"}


def read_rows(path: Path) -> list[tuple[int, dict[str, str]]]:
    """Read normalized worker rows and preserve their source row numbers."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        header_row = None
        headers: list[str] = []
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            normalized = [HEADER_ALIASES.get(header_key(item), header_key(item)) for item in row]
            if REQUIRED_HEADERS.issubset(set(normalized)):
                header_row, headers = row_number, normalized
                break
        if header_row is None:
            raise ValueError("RSR workbook does not contain the expected Worker Details header row.")
        result = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            if not any(item not in (None, "") for item in row):
                continue
            payload = {headers[index]: value(row[index] if index < len(row) else None) for index in range(len(headers)) if headers[index]}
            payload = normalize_payload(payload)
            if is_report_footer(payload, REQUIRED_HEADERS):
                continue
            result.append((row_number, payload))
        return result
    finally:
        workbook.close()
