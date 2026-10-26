"""Source-template-driven CSV/XLSX parsing and structural validation."""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from ..ingestion.date_values import parse_date_value
from .contracts import ParsedSourceRow, SourceTemplate, TemplateConfigurationError


MAX_RUNTIME_FILE_BYTES = 20 * 1024 * 1024
ALLOWED_TRANSFORMS = {"identity", "trim", "lower", "upper", "parse_date", "normalize_email", "normalize_id"}


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _transform(value: str, name: str) -> str:
    if name not in ALLOWED_TRANSFORMS:
        raise TemplateConfigurationError(f"Unsupported runtime transform: {name}")
    if name == "identity":
        return value
    if name in {"trim", "normalize_id", "parse_date"}:
        return value.strip()
    if name == "lower":
        return value.lower()
    if name == "upper":
        return value.upper()
    if name == "normalize_email":
        return value.strip().lower()
    raise TemplateConfigurationError(f"Unsupported runtime transform: {name}")


def _compose_headers(rows: list[tuple[Any, ...]], header_rows: tuple[int, ...]) -> tuple[list[str], int]:
    if not header_rows or any(row < 1 for row in header_rows):
        raise TemplateConfigurationError("Template header rows must be positive row numbers.")
    last_header = max(header_rows)
    if last_header > len(rows):
        raise TemplateConfigurationError("Configured header row is outside the source file.")
    width = max((len(rows[index - 1]) for index in header_rows), default=0)
    headers: list[str] = []
    for column in range(width):
        parts = []
        for row_number in header_rows:
            row = rows[row_number - 1]
            value = _text(row[column]) if column < len(row) else ""
            if value.strip() and value.strip() not in parts:
                parts.append(value.strip())
        headers.append(" · ".join(parts))
    return headers, last_header


def _read_matrix(path: Path, content: bytes, definition: dict[str, Any]) -> tuple[list[tuple[Any, ...]], str]:
    reader = definition.get("reader_settings") or {}
    configured_format = str(reader.get("format", Path(path).suffix.lstrip("."))).upper()
    extension = path.suffix.casefold()
    if extension not in {".csv", ".xlsx"}:
        raise TemplateConfigurationError(f"Unsupported source extension: {extension or '(none)'}")
    if len(content) > MAX_RUNTIME_FILE_BYTES:
        raise TemplateConfigurationError("Runtime source files are limited to 20 MB.")
    if (extension == ".csv" and configured_format != "CSV") or (extension == ".xlsx" and configured_format != "XLSX"):
        raise TemplateConfigurationError("The file extension does not match the approved template reader format.")
    if extension == ".csv":
        delimiter = str(reader.get("csv_delimiter", ",") or ",")
        if len(delimiter) != 1:
            raise TemplateConfigurationError("Configured CSV delimiter must be one character.")
        try:
            content_text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise TemplateConfigurationError("CSV file must be UTF-8 or UTF-8 with BOM.") from exc
        matrix = [tuple(row) for row in csv.reader(io.StringIO(content_text, newline=""), delimiter=delimiter)]
        return matrix, "CSV file"

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise TemplateConfigurationError("XLSX support requires openpyxl.") from exc
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise TemplateConfigurationError(f"Unable to read XLSX workbook: {exc}") from exc
    sheet_name = str(reader.get("sheet_name", "") or "")
    if sheet_name and sheet_name not in workbook.sheetnames:
        workbook.close()
        raise TemplateConfigurationError(f"Configured worksheet {sheet_name!r} was not found.")
    selected_sheet = sheet_name or workbook.sheetnames[0]
    matrix = [tuple(_excel_display_value(cell.value, cell.number_format) for cell in row)
              for row in workbook[selected_sheet].iter_rows()]
    workbook.close()
    return matrix, selected_sheet


def _excel_display_value(value: Any, number_format: str) -> Any:
    """Preserve common zero-padded integer identifiers stored as Excel numbers."""
    if isinstance(value, int) and not isinstance(value, bool) and re.fullmatch(r"0+", str(number_format or "")):
        width = len(number_format)
        if value < 0:
            return "-" + str(abs(value)).zfill(width)
        return str(value).zfill(width)
    return value


def inspect_headers(path: Path, content: bytes, source: SourceTemplate) -> tuple[tuple[str, ...], str]:
    matrix, sheet = _read_matrix(path, content, source.definition)
    reader = source.definition.get("reader_settings") or {}
    header_rows = tuple(int(value) for value in reader.get("header_rows", ()) or ())
    if not header_rows:
        header_rows = (int(reader.get("header_row", 1) or 1),)
    headers, _ = _compose_headers(matrix, header_rows)
    return tuple(headers), sheet


def parse_source_rows(path: Path, content: bytes, source: SourceTemplate) -> tuple[tuple[ParsedSourceRow, ...], tuple[str, ...], tuple[str, ...]]:
    definition = source.definition
    reader = definition.get("reader_settings") or {}
    matrix, sheet = _read_matrix(path, content, definition)
    header_rows = tuple(int(value) for value in reader.get("header_rows", ()) or ())
    if not header_rows:
        header_rows = (int(reader.get("header_row", 1) or 1),)
    headers, last_header_row = _compose_headers(matrix, header_rows)
    if len({header.casefold() for header in headers if header}) != len([header for header in headers if header]):
        raise TemplateConfigurationError("Source file has duplicate composed headers; the configured schema is ambiguous.")

    fields = tuple(field for field in definition.get("fields", ()) if field.get("include", True))
    by_source = {str(field.get("source_name", "")).strip().casefold(): field for field in fields}
    configured_names = {str(field.get("source_name", "")).strip().casefold() for field in fields}
    normalized_headers = {header.casefold() for header in headers}
    unknown_headers = tuple(header for header in headers if header and header.casefold() not in configured_names)
    missing_headers = tuple(
        str(field.get("source_name", "")) for field in fields
        if field.get("required") and str(field.get("source_name", "")).casefold() not in normalized_headers
    )
    issues: list[str] = []
    if missing_headers:
        issues.extend(f"Required source header is missing: {header}" for header in missing_headers)
    if unknown_headers and str(definition.get("schema_drift_policy", "WARNING")).upper() == "BLOCKING":
        issues.extend(f"Blocking schema drift: unconfigured header {header}" for header in unknown_headers)
    if issues:
        return (), tuple(issues), unknown_headers

    rows: list[ParsedSourceRow] = []
    for row_number, values in enumerate(matrix[last_header_row:], start=last_header_row + 1):
        if not any(value not in (None, "") for value in values):
            continue
        raw: dict[str, Any] = {}
        normalized: dict[str, Any] = {}
        row_issues: list[str] = []
        for column, header in enumerate(headers):
            raw_value = values[column] if column < len(values) else ""
            if header:
                raw[header] = _text(raw_value)
            field = by_source.get(header.casefold()) if header else None
            if field is None:
                continue
            key = str(field.get("field_key", "")).strip()
            if not key:
                row_issues.append(f"Configured field {header!r} has no field key.")
                continue
            transform = str(field.get("transform_name", "identity"))
            value = _transform(_text(raw_value), transform)
            if value and str(field.get("data_type", "")).casefold() == "date":
                try:
                    value = parse_date_value(value).isoformat()
                except ValueError:
                    pass
            normalized[key] = value
            row_issues.extend(_validate_field(field, value))
        row_issues.extend(_validate_declared(definition, normalized))
        rows.append(ParsedSourceRow(row_number, sheet, raw, normalized, tuple(dict.fromkeys(row_issues))))
    return tuple(rows), (), unknown_headers


def _validate_field(field: dict[str, Any], value: str) -> list[str]:
    key = str(field.get("field_key", "source field"))
    errors: list[str] = []
    if value == "" and field.get("required"):
        errors.append(f"{key} is required.")
    if value == "" and not field.get("nullable", True):
        errors.append(f"{key} cannot be empty.")
    if value == "":
        return errors
    field_type = str(field.get("data_type", "string")).lower()
    try:
        if field_type == "integer":
            int(value)
        elif field_type == "decimal":
            Decimal(value)
        elif field_type == "boolean" and value.casefold() not in {"true", "false", "yes", "no", "1", "0"}:
            raise ValueError
        elif field_type == "date":
            _parse_date(value)
        elif field_type == "datetime":
            datetime.fromisoformat(value.replace("Z", "+00:00"))
        elif field_type == "enum" and value not in tuple(str(item) for item in field.get("enum_values", ())):
            errors.append(f"{key} is not an allowed value.")
        elif field_type not in {"string", "integer", "decimal", "boolean", "date", "datetime", "enum"}:
            errors.append(f"{key} has unsupported type {field_type!r}.")
    except (ValueError, InvalidOperation, OverflowError):
        errors.append(f"{key} is not a valid {field_type} value.")
    return errors


def _parse_date(value: str) -> date:
    return parse_date_value(value)


def _validate_declared(definition: dict[str, Any], values: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    for rule in definition.get("validations", ()):
        kind = str(rule.get("kind", ""))
        key = str(rule.get("field_key", ""))
        value = str(values.get(key, ""))
        parameter = str(rule.get("parameter", ""))
        if kind == "required" and not value:
            errors.append(f"{key} is required.")
        elif kind == "nullable" and not value:
            errors.append(f"{key} cannot be empty.")
        elif kind == "min_length" and value and len(value) < int(parameter or 0):
            errors.append(f"{key} is shorter than the configured minimum.")
        elif kind == "max_length" and value and len(value) > int(parameter):
            errors.append(f"{key} exceeds the configured maximum length.")
        elif kind == "regex" and value and not re.fullmatch(parameter, value):
            errors.append(f"{key} does not match its configured pattern.")
        elif kind == "enum" and value and value not in [part.strip() for part in parameter.split(",")]:
            errors.append(f"{key} is not an allowed value.")
        elif kind == "date_format" and value:
            try:
                _parse_date(value)
            except ValueError:
                errors.append(f"{key} does not match a supported date format.")
        elif kind not in {"required", "nullable", "min_length", "max_length", "regex", "enum", "date_format", "identifier", "duplicate_source_key"}:
            errors.append(f"Unsupported validation rule {kind!r}.")
    return errors
