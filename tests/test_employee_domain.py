"""
File: tests/test_employee_domain.py
Purpose: Provides automated regression tests for test employee domain.
Responsibilities: Exposes or coordinates: EmployeeDomainTests.
Dependencies: unittest, openpyxl, iam_automation
Inputs/outputs: Inputs: temporary databases, fixtures, and test doubles. Outputs: assertions and test results.
Author / implementation: Miguel Fernandez.
Developed for: Kaiser Permanente NSES GWFM.
Project implementation and integration: Miguel Fernandez.
Third-party source, libraries, and dependencies retain their original licenses.
Estimated implementation effort: 1.25 h (Git-derived heuristic; see docs/file-effort-history.md).
"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from iam_automation.application.employee_service import EmployeeService
from iam_automation.application.roster_lock_ingestion_service import RosterLockIngestionService
from iam_automation.database.runtime_schema import initialize_runtime_database
from iam_automation.infrastructure.sqlite.orm import ImportBatchModel, IngestionRecordModel, WorkerEngagementModel, WorkerProfileVersionModel, orm_session


class EmployeeDomainTests(unittest.TestCase):
    def test_roster_lock_is_staged_before_explicit_worker_promotion(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            database = initialize_runtime_database(root / "runtime.sqlite")
            source = root / "roster.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Vendor Name", "Vendor employee ID", "Employee first name", "Employee last name", "Vendor Email Address", "Job Title", "Program"])
            sheet.append(["TTEC", "E100", "Ada", "Lovelace", "Ada@Example.com", "Analyst", "IAM"])
            workbook.save(source)

            result = RosterLockIngestionService(database).ingest_file(source)
            self.assertEqual(result.staged_rows, 1)
            with orm_session(database) as session:
                record = session.query(IngestionRecordModel).one()
                record_id = record.record_id
            self.assertEqual(EmployeeService(database).list_workers(), [])

            worker_id = EmployeeService(database).promote_ingestion_record(record_id, "admin")
            workers = EmployeeService(database).list_workers()
            self.assertEqual(len(workers), 1)
            self.assertEqual(workers[0].worker_id, worker_id)
            self.assertEqual(workers[0].display_name, "Ada Lovelace")
            self.assertEqual(workers[0].email, "ada@example.com")
            self.assertEqual(workers[0].vendor_name, "TTEC")
            self.assertEqual(workers[0].change_type, "NEW_HIRE")
            roster_workers, roster_total = EmployeeService(database).list_workers_page(source="ROSTER_LOCK", page=1, per_page=25)
            self.assertEqual(roster_total, 1)
            self.assertEqual([item.worker_id for item in roster_workers], [worker_id])
            other_workers, other_total = EmployeeService(database).list_workers_page(source="RSR_SUPPLIER_WORKER_DETAILS", page=1, per_page=25)
            self.assertEqual(other_total, 0)
            self.assertEqual(other_workers, [])
            engagement_workers, engagement_total = EmployeeService(database).list_workers_page(engagement="Analyst", page=1, per_page=25)
            self.assertEqual(engagement_total, 1)
            self.assertEqual([item.worker_id for item in engagement_workers], [worker_id])

    def test_same_vendor_identifier_reuses_worker_and_adds_engagement(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            database = initialize_runtime_database(root / "runtime.sqlite")
            service = EmployeeService(database)
            for index, project in enumerate(("Project A", "Project B"), start=1):
                source = root / f"roster-{index}.xlsx"
                workbook = Workbook()
                sheet = workbook.active
                sheet.append(["Vendor Name", "Vendor employee ID", "Employee first name", "Employee last name", "Project"])
                sheet.append(["TTEC", "E100", "Ada", "Lovelace", project])
                workbook.save(source)
                RosterLockIngestionService(database).ingest_file(source)
                with orm_session(database) as session:
                    record_id = session.query(IngestionRecordModel).order_by(IngestionRecordModel.created_at_utc.desc()).first().record_id
                worker_id = service.promote_ingestion_record(record_id, "admin")
                if index == 1:
                    first_worker = worker_id
            detail = service.get_worker(first_worker)
            self.assertEqual(len(detail["engagements"]), 2)
            roster_profile = next(item for item in detail["source_profiles"] if item["source_system"] == "ROSTER_LOCK")
            self.assertTrue(roster_profile["source_file"].endswith("roster-2.xlsx"))
            self.assertTrue(roster_profile["observed_at_utc"])
            self.assertTrue(roster_profile["engagements"][0]["source_file"])
            self.assertEqual(len(service.list_workers()), 1)

            self.assertEqual(service.list_workers()[0].change_type, "UPDATED")

    def test_reimport_without_material_changes_does_not_create_engagement_or_version(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            database = initialize_runtime_database(root / "runtime.sqlite")
            source = root / "roster.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Vendor Name", "Vendor employee ID", "Employee first name", "Employee last name", "Job Title", "Program"])
            sheet.append(["TTEC", "E100", "Ada", "Lovelace", "Analyst", "IAM"])
            workbook.save(source)
            ingestion = RosterLockIngestionService(database)
            ingestion.ingest_file(source)
            service = EmployeeService(database)
            with orm_session(database) as session:
                first = session.query(IngestionRecordModel).one()
                worker_id = service.promote_ingestion_record(first.record_id, "admin")
                before_engagements = session.query(WorkerEngagementModel).filter_by(worker_id=worker_id).count()
                before_versions = session.query(WorkerProfileVersionModel).filter_by(worker_id=worker_id).count()
                session.add(ImportBatchModel(batch_id="reimport-batch", filename="roster-copy.xlsx", source_type="ROSTER_LOCK", status="COMPLETED", total_rows=1, valid_rows=1, created_by="admin"))
                session.add(IngestionRecordModel(record_id="reimport-record", batch_id="reimport-batch", source_system="ROSTER_LOCK", source_file="/source/roster-copy.xlsx", source_checksum="copy-checksum", row_number=2, source_record_key="E100", normalized_json='{"job_title":"Analyst","program":"IAM","vendor_employee_id":"E100","first_name":"Ada","last_name":"Lovelace","vendor_name":"TTEC"}', payload_json="", status="STAGED", created_at_utc="2026-09-07T12:00:00+00:00"))

            result = service.promote_staged_records(["reimport-record"], "admin")

            self.assertEqual(result["processed"], 1)
            self.assertEqual(result["created"], 0)
            self.assertEqual(result["reused"], 1)
            self.assertEqual(result["changed"], 0)
            self.assertEqual(result["unchanged"], 1)
            with orm_session(database) as session:
                self.assertEqual(session.query(WorkerEngagementModel).filter_by(worker_id=worker_id).count(), before_engagements)
                self.assertEqual(session.query(WorkerProfileVersionModel).filter_by(worker_id=worker_id).count(), before_versions)

    def test_directory_change_badge_classifies_termination_and_missing_history(self):
        self.assertEqual(EmployeeService._directory_change_type("ACTIVE", "Terminated", "ROSTER_LOCK_ENRICHMENT"), "TERMINATION")
        self.assertEqual(EmployeeService._directory_change_type("ACTIVE", "New Hire", "ROSTER_LOCK_INITIAL_PROFILE"), "NEW_HIRE")
        self.assertEqual(EmployeeService._directory_change_type("ACTIVE", "", ""), "NO_CHANGE")

    def test_bulk_promotion_processes_staged_rows_without_one_click_per_worker(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            database = initialize_runtime_database(root / "runtime.sqlite")
            source = root / "roster-bulk.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Vendor Name", "Vendor employee ID", "Employee first name", "Employee last name", "Vendor Email Address", "Project"])
            sheet.append(["TTEC", "E100", "Ada", "Lovelace", "ada@example.com", "Project A"])
            sheet.append(["TTEC", "E101", "Grace", "Hopper", "grace@example.com", "Project A"])
            sheet.append(["TTEC", "E102", "Katherine", "Johnson", "katherine@example.com", "Project B"])
            workbook.save(source)

            result = RosterLockIngestionService(database).ingest_file(source)
            self.assertEqual(result.staged_rows, 3)

            promotion = EmployeeService(database).promote_all_staged_records("admin")

            self.assertEqual(promotion["processed"], 3)
            self.assertEqual(promotion["created"], 3)
            self.assertEqual(promotion["reused"], 0)
            self.assertEqual(promotion["failed"], 0)
            self.assertEqual(len(EmployeeService(database).list_workers()), 3)
            with orm_session(database) as session:
                self.assertEqual(session.query(IngestionRecordModel).filter_by(status="MATCHED").count(), 3)

    def test_source_update_preview_returns_only_real_changes_for_existing_worker(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            database = initialize_runtime_database(root / "runtime.sqlite")
            service = EmployeeService(database)

            original = root / "original.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Vendor Name", "Vendor employee ID", "Employee first name", "Employee last name", "Job Title"])
            sheet.append(["TTEC", "E100", "Ada", "Lovelace", "Analyst"])
            workbook.save(original)
            RosterLockIngestionService(database).ingest_file(original)
            with orm_session(database) as session:
                record_id = session.query(IngestionRecordModel).one().record_id
            worker_id = service.promote_ingestion_record(record_id, "admin")

            changed = root / "changed.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Vendor Name", "Vendor employee ID", "Employee first name", "Employee last name", "Job Title"])
            sheet.append(["TTEC", "E100", "Ada", "Lovelace", "Senior Analyst"])
            workbook.save(changed)
            RosterLockIngestionService(database).ingest_file(changed)
            with orm_session(database) as session:
                staged = session.query(IngestionRecordModel).filter_by(status="STAGED").one()
                staged.matched_worker_id = worker_id
                staged.match_status = "MATCHED"
                staged.match_reason = "Exact identifier match"

            preview = service.source_update_preview(worker_id)
            self.assertIsNotNone(preview)
            self.assertEqual(len(preview["updates"]), 1)
            changes = preview["updates"][0]["changes"]
            self.assertEqual([(item["field"], item["current_value"], item["incoming_value"]) for item in changes], [("job_title", "Analyst", "Senior Analyst")])
            fields = {item["field"]: item for item in preview["fields"]}
            self.assertTrue(fields["job_title"]["changed"])
            self.assertEqual(fields["job_title"]["current_value"], "Analyst")
            self.assertEqual(fields["job_title"]["incoming_value"], "Senior Analyst")
            self.assertFalse(fields["vendor_name"]["changed"])


if __name__ == "__main__":
    unittest.main()
#endfile
