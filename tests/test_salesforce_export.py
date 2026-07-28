from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from iam_automation.application.request_service import RequestService
from iam_automation.application.salesforce_export_service import SalesforceExportService
from iam_automation.database.runtime_schema import initialize_runtime_database
from iam_automation.identity import CurrentUser, ROLE_REQUESTER, ROLE_REVIEWER


class SalesforceExportTests(unittest.TestCase):
    def test_generates_official_template_from_approved_sqlite_request(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            database = initialize_runtime_database(root / "runtime.sqlite")
            requester = CurrentUser("requester", "VDI-01", (ROLE_REQUESTER,))
            reviewer = CurrentUser("reviewer", "VDI-02", (ROLE_REVIEWER,))
            request = RequestService(database).create(requester, "CREATE_EMPLOYEE", "E100", {"first_name": "Ana", "last_name": "Garcia", "email": "ana.garcia@example.org", "lob": "NCS"}, "Approved test export")
            RequestService(database).decide(request.request_id, reviewer, "APPROVED", "Approved for Salesforce export")

            destination = SalesforceExportService(database).generate(root / "generated_templates", "master")
            worksheet = load_workbook(destination, data_only=False).active
            headers = {cell.value: cell.column for cell in worksheet[2] if cell.value}
            self.assertEqual(worksheet.cell(3, headers["FIRSTNAME"]).value, "Ana")
            self.assertEqual(worksheet.cell(3, headers["LASTNAME"]).value, "Garcia")
            self.assertEqual(worksheet.cell(3, headers["USERNAME"]).value, "ana.garcia@example.org")
            self.assertEqual(worksheet.cell(3, headers["EMPLOYEENUMBER"]).value, "E100")
            self.assertTrue(str(worksheet.cell(3, headers["NAME"]).value).startswith("=CONCATENATE"))
            self.assertNotEqual(worksheet.cell(3, headers["FIRSTNAME"]).value, "Mercille")

    def test_requires_approved_data(self):
        with tempfile.TemporaryDirectory() as tmp:
            database = initialize_runtime_database(Path(tmp) / "runtime.sqlite")
            with self.assertRaisesRegex(ValueError, "No approved"):
                SalesforceExportService(database).generate(Path(tmp) / "generated_templates", "master")


if __name__ == "__main__":
    unittest.main()
