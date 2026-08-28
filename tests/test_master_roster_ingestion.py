from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from sqlalchemy import select

from iam_automation.application.master_roster_ingestion_service import MasterRosterIngestionService
from iam_automation.application.employee_service import EmployeeService
from iam_automation.application.file_watch_service import FileWatchService
from iam_automation.database.runtime_schema import initialize_runtime_database
from iam_automation.application.notification_service import NotificationService
from iam_automation.infrastructure.sqlite.orm import ChangeRequestModel, ImportBatchModel, IngestionRecordModel, orm_session


class MasterRosterIngestionTests(unittest.TestCase):
    def _write_workbook(self, path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([None] * 18 + ["System Type Name", "CLIENT", "OPERATIONS"])
        sheet.append([None] * 18 + ["System Name", "Kaiser Systems", "Healthcare Systems"])
        sheet.append([None] * 18 + ["Sub System Name", "NUID", "NPN"])
        sheet.append([None, "Parent Client", "Client", "Program", "Project", "Site Hr Location", "Employee Number", "Employee Name", "Supervisor Employee Number", "Supervisor Name", "Job Code", "Job Title", "Hire Date", "Project Start Date", "Production Start Date", "Actual Termination Date", "Term Reason", "Assignment Status", "User Principal Name", "LogonID", "LogonID"])
        sheet.append([1, "Kaiser", "Client", "Program", "Project", "Site", "E-100", "Ada Lovelace", "S-1", "Supervisor", "J-1", "Analyst", "2020-01-01", "2025-01-01", "2025-01-01", None, None, "Active Assignment", "ada@example.com", "NUID-100", "NPN-100"])
        sheet.append([2, "Kaiser", "Client", "Program", "Project", "Site", "E-101", "Grace Hopper", "S-2", "Supervisor", "J-2", "Engineer", "2020-01-01", "2025-01-01", "2025-01-01", None, None, "Active Assignment", None, "NUID-101", "NPN-101"])
        workbook.save(path)

    def test_stages_metadata_header_and_both_login_ids_without_canonical_mutation(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            database = initialize_runtime_database(root / "runtime.sqlite")
            workbook = root / "master.xlsx"
            self._write_workbook(workbook)

            result = MasterRosterIngestionService(database).ingest_file(workbook, actor="test")

            self.assertEqual(result.total_rows, 2)
            self.assertEqual(result.staged_rows, 2)
            self.assertEqual(result.quarantine_rows, 0)
            with orm_session(database) as session:
                batch = session.get(ImportBatchModel, result.batch_id)
                record = session.scalar(select(IngestionRecordModel).where(IngestionRecordModel.source_record_key == "E-100"))
                self.assertEqual(batch.source_type, "MASTER_ROSTER")
                self.assertEqual(record.row_number, 5)
                self.assertIn("nuid_login_id", record.normalized_json)
                self.assertIn("NPN-100", record.normalized_json)
                self.assertEqual(session.query(IngestionRecordModel).count(), 2)

    def test_same_file_is_idempotent(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            database = initialize_runtime_database(root / "runtime.sqlite")
            workbook = root / "master.xlsx"
            self._write_workbook(workbook)
            service = MasterRosterIngestionService(database)

            first = service.ingest_file(workbook)
            second = service.ingest_file(workbook)

            self.assertFalse(first.already_processed)
            self.assertTrue(second.already_processed)
            self.assertEqual(first.batch_id, second.batch_id)

    def test_manual_import_promotes_all_staged_rows_to_history_and_creates_workers(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            database = initialize_runtime_database(root / "runtime.sqlite")
            workbook = root / "master.xlsx"
            self._write_workbook(workbook)
            ingestion = MasterRosterIngestionService(database)
            result = ingestion.ingest_file(workbook, actor="test")
            with orm_session(database) as session:
                record_ids = list(session.scalars(select(IngestionRecordModel.record_id).where(IngestionRecordModel.batch_id == result.batch_id)).all())

            applied = EmployeeService(database).apply_approved_import_records(record_ids, "test")

            self.assertEqual(applied["processed"], 2)
            self.assertEqual(applied["created"], 2)
            with orm_session(database) as session:
                self.assertEqual(session.query(IngestionRecordModel).filter_by(status="MATCHED").count(), 2)

    def test_duplicate_employee_number_is_quarantined(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            database = initialize_runtime_database(root / "runtime.sqlite")
            workbook = root / "master.xlsx"
            self._write_workbook(workbook)
            from openpyxl import load_workbook
            editable = load_workbook(workbook)
            editable.active.cell(row=6, column=7).value = "E-100"
            editable.save(workbook)

            result = MasterRosterIngestionService(database).ingest_file(workbook)

            self.assertEqual(result.staged_rows, 1)
            self.assertEqual(result.quarantine_rows, 1)

    def test_watcher_stages_master_roster_without_creating_approval_request(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            database = initialize_runtime_database(root / "runtime.sqlite")
            source_folder = root / "master_roster"
            source_folder.mkdir()
            self._write_workbook(source_folder / "master.xlsx")

            processed = FileWatchService(database, root).scan_once()

            self.assertEqual(processed, 1)
            with orm_session(database) as session:
                self.assertEqual(session.query(IngestionRecordModel).filter_by(source_system="MASTER_ROSTER").count(), 2)
                self.assertEqual(session.query(ChangeRequestModel).count(), 0)

    def test_watcher_creates_master_roster_notification_without_cycle_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            database = initialize_runtime_database(root / "runtime.sqlite")
            source_folder = root / "master_roster"
            source_folder.mkdir()
            self._write_workbook(source_folder / "master.xlsx")
            service = FileWatchService(database, root)
            service._active_admins = lambda: ("admin",)

            self.assertEqual(service.scan_once(), 1)
            notification = next(item for item in NotificationService(database).list_for_user("admin") if item.notification_type == "FILE_DETECTED")
            self.assertEqual(notification.title, "Master Roster ingested")


if __name__ == "__main__":
    unittest.main()
